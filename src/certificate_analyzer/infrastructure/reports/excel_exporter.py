from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from certificate_analyzer.infrastructure.reports.rows import (
    HEADERS,
    model_rows,
    values,
    output_path,
)


class ExcelReportExporter:
    def __init__(self, export_folder=""):
        self.export_folder = export_folder

    def export(
        self, certificates, mchds, save_path=None, report_title="Отчет по данным"
    ):
        return self.export_rows(
            model_rows(certificates, mchds), save_path, report_title
        )

    def export_rows(self, rows, save_path=None, report_title="Отчет"):
        target = output_path(self.export_folder, save_path, "xlsx")
        book = Workbook()
        sheet = book.active
        sheet.title = "Данные"
        sheet.append([report_title])
        sheet.merge_cells("A1:K1")
        sheet.append([])
        sheet.append([])
        for col, header in enumerate(HEADERS, 1):
            cell = sheet.cell(4, col, header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="6C5CE7")
        for row in rows:
            sheet.append(values(row))
            status = row.get("status", "")
            color = (
                "F5C6CB"
                if status == "Просрочен"
                else "FFEEBA"
                if "Истекает" in status
                else "C3E6CB"
            )
            for cell in sheet[sheet.max_row]:
                # Imported personal data must stay text, including values starting with '='.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.fill = PatternFill("solid", fgColor=color)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet["A1"].data_type = "s"
        sheet.freeze_panes = "A5"
        sheet.auto_filter.ref = f"A4:K{sheet.max_row}"
        for index in range(1, len(HEADERS) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(
                50,
                max(
                    12,
                    max(
                        len(str(sheet.cell(r, index).value or ""))
                        for r in range(4, sheet.max_row + 1)
                    )
                    + 2,
                ),
            )
        book.save(target)
        return target
